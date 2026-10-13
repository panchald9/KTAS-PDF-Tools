"""Fast local engine regression checks; synthetic documents only, no external APIs."""
import base64
import contextlib
import io
import json
import sys
import tempfile
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root / "artifacts/api-server/python"))
from worker import process


with tempfile.TemporaryDirectory(prefix="ktas-check-") as temporary:
    work = Path(temporary)
    pdf = fitz.open()
    for i in range(3):
        p = pdf.new_page()
        p.insert_text((50, 60), f"Page {i+1}. SECRET CLIENT DATA. This is an English document about solar energy.")
        p.insert_text((50, 90), "Solar power reduces electricity costs. Energy storage helps during the night.")
        p.insert_text((50, 120), "Panels must be maintained. Clean power is useful for homes and businesses.")
    widget = fitz.Widget()
    widget.field_name, widget.field_type = "customer", fitz.PDF_WIDGET_TYPE_TEXT
    widget.rect, widget.field_value = fitz.Rect(50, 160, 250, 185), "Original"
    pdf[0].add_widget(widget)
    pdf.set_metadata({"subject": "SECRET CLIENT DATA"})
    source = work / "source.pdf"
    pdf.save(source)
    other = work / "other.pdf"
    pdf[1].insert_text((50, 220), "Changed text in the second version.")
    pdf.save(other)
    pdf.close()
    large = fitz.open()
    for i in range(82):
        large.new_page().insert_text((30, 50), f"Actual page {i+1}")
    large_path = work / "82pages.pdf"
    large.save(large_path)
    large.close()
    image = work / "image.png"
    Image.new("RGB", (480, 640), "white").save(image)
    scan = fitz.open()
    p = scan.new_page()
    p.insert_text((50, 80), "LOCAL OCR CHECK DOCUMENT", fontsize=22)
    p.insert_text((50, 120), "The quick brown fox jumps over the lazy dog.", fontsize=16)
    pix = p.get_pixmap(dpi=150)
    pix.save(work / "scan.png")
    scan.close()
    word = Document()
    word.add_paragraph("This is a real Word source document.")
    word_path = work / "word.docx"
    word.save(word_path)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Real PowerPoint slide"
    deck_path = work / "slides.pptx"
    deck.save(deck_path)
    book = Workbook()
    book.active.append(["Item", "Amount"])
    book.active.append(["Example", 42])
    book_path = work / "book.xlsx"
    book.save(book_path)
    html_path = work / "page.html"
    html_path.write_text("<h1>Actual uploaded HTML</h1><p>Rendered document content.</p><script>alert(1)</script>")
    cases = {
        "merge-pdf": ([source, other], {}),
        "split-pdf": ([source], {"pages":"2-3"}),
        "remove-pages": ([source], {"pages":"2"}),
        "organize-pdf": ([large_path], {"order":""}),
        "scan-to-pdf": ([image], {}),
        "compress-pdf": ([source], {"quality":"balanced"}),
        "repair-pdf": ([source], {}),
        "ocr-pdf": ([work / "scan.png"], {}),
        "jpg-to-pdf": ([image, image], {}),
        "word-to-pdf": ([word_path], {}),
        "powerpoint-to-pdf": ([deck_path], {}),
        "excel-to-pdf": ([book_path], {}),
        "html-to-pdf": ([html_path], {}),
        "pdf-to-jpg": ([source], {"pages":"1,3","dpi":"72"}),
        "pdf-to-word": ([source], {}),
        "pdf-to-powerpoint": ([source], {}),
        "pdf-to-excel": ([source], {}),
        "pdf-to-pdfa": ([source], {}),
        "rotate-pdf": ([source], {"angle":"90","pages":"2"}),
        "page-numbers": ([source], {"start":"8"}),
        "watermark-pdf": ([source], {"text":"ACTUAL WATERMARK","opacity":"0.2"}),
        "crop-pdf": ([source], {"margin":"20"}),
        "edit-pdf": ([source], {"text":"ACTUAL ADDED TEXT","page":"2","x":"40","y":"250"}),
        "pdf-forms": ([source], {"fields":json.dumps({"customer":"Filled customer"})}),
        "protect-pdf": ([source], {"password":"Synthetic-secret-42"}),
        "sign-pdf": ([source], {"signatureText":"Synthetic Signature","page":"1","x":"40","y":"250"}),
        "redact-pdf": ([source], {"redactText":"SECRET CLIENT DATA","rectangles":"[]"}),
        "compare-pdf": ([source, other], {}),
        "ai-summarizer": ([source], {"sentences":"3"}),
        "translate-pdf": ([source], {"language":"es"}),
        "pdf-to-markdown": ([source], {}),
    }
    failed, passed, outputs = [], [], {}
    for slug, (paths, options) in cases.items():
        # Engines can create private subfolders: a separate work directory for every request.
        case_dir = work / slug
        case_dir.mkdir()
        local = []
        for i, path in enumerate(paths):
            p = case_dir / f"input-{i}{path.suffix}"
            p.write_bytes(path.read_bytes())
            local.append(str(p))
        try:
            with contextlib.redirect_stdout(sys.stderr):
                result = process({"slug":slug,"paths":local,"options":options})
            file = result["files"][0]
            data = base64.b64decode(file["dataBase64"])
            assert len(data) == file["bytes"] == result["outputBytes"]
            if file["name"].endswith(".pdf"):
                doc = fitz.open(stream=data, filetype="pdf")
                if slug == "protect-pdf":
                    assert doc.needs_pass and not doc.authenticate("wrong")
                    assert doc.authenticate("Synthetic-secret-42")
                if slug == "organize-pdf":
                    assert len(doc) == 82 and "Actual page 82" in doc[81].get_text()
                if slug == "merge-pdf":
                    assert len(doc) == 6 and any(list(p.widgets() or []) for p in doc)
                if slug == "redact-pdf":
                    assert "SECRET CLIENT DATA" not in "".join(p.get_text() for p in doc)
                    assert "SECRET CLIENT DATA" not in str(doc.metadata)
                if slug == "pdf-forms":
                    assert list(doc[0].widgets())[0].field_value == "Filled customer"
                if slug == "rotate-pdf":
                    assert doc[0].rotation == 0 and doc[1].rotation == 90 and len(doc) == 3
                if slug == "edit-pdf":
                    assert "ACTUAL ADDED TEXT" in doc[1].get_text()
                if slug == "page-numbers":
                    assert "8" in doc[0].get_text() and "10" in doc[2].get_text()
                if slug == "watermark-pdf":
                    assert "ACTUAL WATERMARK" in doc[0].get_text()
                if slug == "ocr-pdf":
                    assert "OCR" in doc[0].get_text()
                doc.close()
            elif file["name"].endswith((".zip", ".docx", ".pptx", ".xlsx")):
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    names = z.namelist()
                    if slug == "pdf-to-jpg":
                        assert len(names) == 2
                        for name in names:
                            with Image.open(io.BytesIO(z.read(name))) as check:
                                assert check.format == "JPEG"
                    if slug == "pdf-to-word": assert "word/document.xml" in names
                    if slug == "pdf-to-powerpoint": assert "ppt/presentation.xml" in names
                    if slug == "pdf-to-excel": assert "xl/workbook.xml" in names
            if slug == "ai-summarizer": assert "solar" in result["text"].lower()
            if slug == "translate-pdf": assert result["text"] and "energ" in result["text"].lower()
            outputs[slug] = data
            passed.append(slug)
            print("PASS", slug, len(data), flush=True)
        except Exception as error:
            failed.append((slug, str(error)))
            import traceback
            traceback.print_exc()
            print("FAIL", slug, str(error), flush=True)
    if "protect-pdf" in outputs:
        path = work / "locked.pdf"
        path.write_bytes(outputs["protect-pdf"])
        result = process({"slug":"unlock-pdf","paths":[str(path)],"options":{"password":"Synthetic-secret-42"}})
        doc = fitz.open(stream=base64.b64decode(result["files"][0]["dataBase64"]),filetype="pdf")
        assert not doc.needs_pass and len(doc) == 3
        print("PASS unlock-pdf", flush=True)
        passed.append("unlock-pdf")
    print(json.dumps({"passed":len(passed),"failed":failed}), flush=True)
    if failed:
        raise SystemExit(1)