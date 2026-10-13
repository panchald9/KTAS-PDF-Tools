"""Validate actual Word outputs, reopened DOCX content and malformed input handling."""
import base64
import io
import json
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "artifacts/api-server/python"))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from word_tools import dispatch, WORD_SLUGS, packed, parse


def decoded(result):
    assert result["files"] and result["outputBytes"] > 0
    return base64.b64decode(result["files"][0]["dataBase64"])


def paragraphs(doc):
    return "\n".join(p.text for p in doc.paragraphs)


with tempfile.TemporaryDirectory(prefix="ktas-word-check-") as temporary:
    root = Path(temporary)
    source = root / "source.docx"
    second = root / "second.docx"
    plain = root / "source.txt"
    plain.write_text("Hello world\nUnicode café नमस्ते", encoding="utf-8")
    doc = Document()
    doc.core_properties.author = "Test author"
    p = doc.add_paragraph()
    p.add_run("Hello ")
    p.add_run("wor").bold = True
    p.add_run("ld and world.")
    doc.add_paragraph("<script>alert('no')</script>")
    doc.add_paragraph()
    doc.add_paragraph("   ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "=1+1"
    table.cell(0, 1).text = "Header"
    table.cell(1, 0).text = "café"
    table.cell(1, 1).text = "नमस्ते"
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "red").save(image, format="PNG")
    image.seek(0)
    doc.add_picture(image)
    doc.add_paragraph().add_run().add_break()
    doc.add_paragraph("Change ")
    p = doc.paragraphs[-1]._p
    ins = OxmlElement("w:ins")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "inserted"
    run.append(text)
    ins.append(run)
    p.append(ins)
    deletion = OxmlElement("w:del")
    run, text = OxmlElement("w:r"), OxmlElement("w:delText")
    text.text = "deleted"
    run.append(text)
    deletion.append(run)
    p.append(deletion)
    # A real comment with anchors and relationship.
    doc.add_comment(doc.paragraphs[0].runs, "Private comment", author="Reviewer")
    doc.sections[0].header.paragraphs[0].text = "Header world"
    doc.save(source)
    extra = Document()
    extra.add_paragraph("SECOND DOCUMENT")
    extra.save(second)
    for slug in sorted(WORD_SLUGS):
        paths = [plain] if slug == "text-to-word" else [source, second] if slug == "merge-word" else [source]
        opts = {"find": "world", "replace": "planet", "header": "New header", "footer": "New footer"}
        result = dispatch(slug, paths, opts)
        data = decoded(result)
        if result["files"][0]["name"].endswith(".docx"):
            reopened = Document(io.BytesIO(data))
            text = paragraphs(reopened)
            if slug == "text-to-word":
                assert "नमस्ते" in text and "Hello world" in text
            elif slug == "merge-word":
                assert "SECOND DOCUMENT" in text and "Hello world" in text
                assert len(reopened.inline_shapes) == 1
            elif slug == "find-replace-word":
                assert "Hello planet and planet." in text and "world" not in text
                assert "planet" in reopened.sections[0].header.paragraphs[0].text
                assert len(reopened.inline_shapes) == 1
            elif slug == "accept-word-revisions":
                assert "inserted" in text and "deleted" not in text
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    xml = archive.read("word/document.xml")
                    assert b"<w:ins " not in xml and b"<w:del " not in xml
            elif slug == "remove-word-comments":
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    assert not any("comments" in name.lower() for name in archive.namelist())
                    assert b"commentReference" not in archive.read("word/document.xml")
                    assert b"/comments" not in archive.read("word/_rels/document.xml.rels")
            elif slug == "remove-word-metadata":
                assert not reopened.core_properties.author
                assert not reopened.core_properties.created
            elif slug == "word-page-size":
                assert round(reopened.sections[0].page_width.mm) == 210
            elif slug == "word-margins":
                assert round(reopened.sections[0].left_margin.mm) == 25
            elif slug == "word-font":
                assert reopened.paragraphs[0].runs[0].font.name == "Arial"
                assert reopened.paragraphs[0].runs[0].font.size.pt == 12
            elif slug == "word-line-spacing":
                assert reopened.paragraphs[0].paragraph_format.line_spacing == 1.5
            elif slug == "word-header-footer":
                assert "New header" in reopened.sections[0].header.paragraphs[0].text
            elif slug == "remove-word-empty-paragraphs":
                assert len(reopened.inline_shapes) == 1 and "<w:br" in reopened.element.xml
                assert len(reopened.paragraphs) < len(doc.paragraphs)
        elif slug == "word-to-text":
            assert "Hello world" in data.decode() and "नमस्ते" in data.decode()
        elif slug == "word-to-html":
            assert b"&lt;script&gt;" in data and b"<script>" not in data
        elif slug == "word-statistics":
            assert json.loads(data)["topLevelTables"] == 1 and json.loads(data)["pageCount"] is None
        elif slug == "word-metadata":
            assert json.loads(data)["author"] == "Test author"
        elif slug == "word-tables-to-csv":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                assert "'=1+1" in archive.read("table-1.csv").decode("utf-8-sig")
        elif slug == "extract-word-images":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                assert archive.read(archive.namelist()[0]).startswith(b"\x89PNG")
        print("PASS", slug)
    for slug, opts, paths in [
        ("find-replace-word", {"find": ""}, [source]),
        ("word-font", {"fontSize": "nan"}, [source]),
        ("word-line-spacing", {"spacing": "99"}, [source]),
        ("word-page-size", {"size": "unsupported"}, [source]),
        ("word-margins", {"left": "101"}, [source]),
        ("merge-word", {}, [source]),
    ]:
        try:
            dispatch(slug, paths, opts)
            raise AssertionError(f"{slug} accepted invalid options")
        except ValueError:
            print("PASS invalid", slug)
    # Reject malformed Office packages and dangerous XML, without extracting files.
    for index, parts in enumerate([
        {"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<doc/>", "../outside": b"x"},
        {"[Content_Types].xml": b"<Types/>", "word/document.xml": b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>'},
        {"[Content_Types].xml": b"<Types/>"},
    ]):
        bad = root / f"bad-{index}.docx"
        bad.write_bytes(packed(parts))
        try:
            dispatch("word-to-text", [bad], {})
            raise AssertionError("Accepted invalid package")
        except ValueError:
            print("PASS invalid package", index)
print("All Word tool checks passed.")