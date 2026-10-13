"""Bounded local XLSX/CSV tools. No formula evaluation or macro execution."""
import csv
import io
import json
import re
import zipfile
from bisect import bisect_left
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import PurePosixPath

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import quote_sheetname
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.xml.constants import XLSX, XLSM, XLTX, XLTM

from common import output, zip_outputs

SLUGS = {
    "merge-excel", "split-excel", "excel-to-csv", "csv-to-excel",
    "remove-excel-duplicates", "remove-excel-empty-rows", "trim-excel-text",
    "excel-statistics", "excel-to-json", "inspect-workbook",
}
MAX_CELLS = 500_000
MAX_RECTANGLE = 500_000
MAX_SHEETS = 100
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
WARNINGS = [
    "Formulas are never evaluated. Existing formula caches are cleared when saving XLSX; recalculate in a trusted spreadsheet editor.",
    "Keep the original. Review formatting and formulas before relying on the result.",
]


def package_check(path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = [e.filename for e in entries]
            if len(entries) > 4000 or sum(e.file_size for e in entries) > 120 * 1024 * 1024:
                raise ValueError("Workbook package exceeds the expanded size limit.")
            if len(names) != len(set(names)):
                raise ValueError("Workbook contains duplicate package parts.")
            if not {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}.issubset(names):
                raise ValueError("Choose a real XLSX workbook, not another ZIP or Office file.")
            cells = merged_cells = 0
            workbook = relationships = manifest = None
            for entry in entries:
                name = PurePosixPath(entry.filename)
                if name.is_absolute() or ".." in name.parts or "\\" in entry.filename or entry.flag_bits & 1:
                    raise ValueError("Invalid or encrypted workbook package.")
                if "vba" in entry.filename.lower() or entry.filename.startswith(("xl/externalLinks/", "xl/macrosheets/")):
                    raise ValueError("Macros and external workbook links are unsupported. Choose a plain XLSX.")
                if entry.filename.endswith((".xml", ".rels")):
                    data = archive.read(entry)
                    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                        raise ValueError("Workbook contains unsupported XML declarations.")
                    root = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
                    if entry.filename == "[Content_Types].xml":
                        manifest = root
                    elif entry.filename == "xl/workbook.xml":
                        workbook = root
                    elif entry.filename == "xl/_rels/workbook.xml.rels":
                        relationships = root
                    if entry.filename == "[Content_Types].xml" and b"macroEnabled" in data:
                        raise ValueError("Macro-enabled workbooks are not accepted.")
                    if entry.filename.startswith("xl/worksheets/"):
                        cells += len(root.findall(".//s:c", NS))
                        if cells + merged_cells > MAX_CELLS:
                            raise ValueError(f"Workbook stored cells and merged-cell expansion exceed {MAX_CELLS:,} cells.")
                        # openpyxl allocates every position in merged ranges while
                        # loading, before read_book can inspect worksheet dimensions.
                        for merged in root.findall(".//s:mergeCell", NS):
                            try:
                                low_col, low_row, high_col, high_row = range_boundaries(merged.get("ref", ""))
                            except (ValueError, TypeError):
                                raise ValueError("Workbook contains an invalid merged-cell range.")
                            if any(value is None for value in (low_col, low_row, high_col, high_row)):
                                raise ValueError("Merged-cell ranges must use finite cell coordinates.")
                            if not (1 <= low_row <= high_row <= 100_000 and 1 <= low_col <= high_col <= 2000):
                                raise ValueError("Merged-cell range exceeds 100,000 rows or 2,000 columns.")
                            merged_cells += (high_row - low_row + 1) * (high_col - low_col + 1)
                            if merged_cells + cells > MAX_CELLS:
                                raise ValueError("Merged-cell expansion exceeds the 500,000 cell budget.")
            # The loader selects the workbook through the manifest, and sheets
            # through relationships, not through ZIP filenames. Only allow the
            # canonical locations whose XML was budgeted above. Validate every
            # sheet's relationship, regardless of its declared relationship type.
            workbook_parts = [node for node in manifest if node.get("ContentType") in {XLSX, XLSM, XLTX, XLTM}]
            if (len(workbook_parts) != 1 or workbook_parts[0].get("ContentType") != XLSX
                    or workbook_parts[0].get("PartName") != "/xl/workbook.xml"
                    or etree.QName(workbook_parts[0]).localname != "Override"):
                raise ValueError("Only standard XLSX workbook package locations are supported.")
            rels = relationships.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
            rel_ids = [rel.get("Id") for rel in rels]
            if len(rel_ids) != len(set(rel_ids)):
                raise ValueError("Workbook contains duplicate relationship IDs.")
            by_id = {rel.get("Id"): rel for rel in rels}
            workbook_sheets = workbook.findall(".//s:sheet", NS)
            if not 1 <= len(workbook_sheets) <= MAX_SHEETS:
                raise ValueError(f"Choose a workbook with 1–{MAX_SHEETS} worksheets.")
            used_targets = set()
            for sheet in workbook_sheets:
                rel = by_id.get(sheet.get(f"{{{OFFICE_REL_NS}}}id"))
                if rel is None:
                    raise ValueError("Workbook has a missing worksheet relationship.")
                target = rel.get("Target", "")
                if (rel.get("TargetMode") == "External"
                        or rel.get("Type") != OFFICE_REL_NS + "/worksheet"
                        or not re.fullmatch(r"(?:/xl/)?worksheets/sheet[1-9]\d*\.xml", target)):
                    raise ValueError("Unsupported worksheet relationship location. Use a standard XLSX workbook.")
                actual = target[1:] if target.startswith("/") else "xl/" + target
                if actual not in names:
                    raise ValueError("Workbook has a missing worksheet part.")
                if actual in used_targets:
                    raise ValueError("Workbook references the same worksheet part more than once.")
                used_targets.add(actual)
    except (zipfile.BadZipFile, etree.XMLSyntaxError, KeyError, RuntimeError):
        raise ValueError("Malformed or encrypted XLSX package.")
    return cells + merged_cells


def read_book(path, cached=False, checked=False):
    if not checked:
        package_check(path)
    try:
        wb = load_workbook(path, data_only=cached, keep_links=False)
    except Exception:
        raise ValueError("This XLSX workbook cannot be read.")
    if not wb.worksheets or len(wb.worksheets) > MAX_SHEETS:
        raise ValueError(f"Choose a workbook with 1–{MAX_SHEETS} worksheets.")
    total = 0
    for ws in wb:
        if ws.max_row > 100_000 or ws.max_column > 2000:
            raise ValueError("Worksheet exceeds 100,000 rows or 2,000 columns.")
        total += ws.max_row * ws.max_column
    if total > MAX_RECTANGLE:
        raise ValueError("Combined worksheet dimensions exceed 500,000 cells. Remove distant formatting and try again.")
    return wb


def selected(wb, options):
    name = options.get("sheet", "")
    if name:
        if name not in wb.sheetnames:
            raise ValueError("The selected worksheet name does not exist (names are case-sensitive).")
        return [wb[name]]
    return wb.worksheets


def delimiter(options):
    values = {"comma": ",", "semicolon": ";", "tab": "\t"}
    value = options.get("delimiter", "comma")
    if value not in values:
        raise ValueError("Choose comma, semicolon or tab as the CSV delimiter.")
    return values[value]


def plain(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return value


def safe_csv(value):
    text = "" if value is None else str(plain(value))
    # Leading controls/whitespace cannot hide a formula trigger.
    if text.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@", "＝", "＋", "－", "＠")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def inspect_workbook(path, options):
    """Read-only sample; bounded independently of the worker output budget."""
    wb = read_book(path)
    cached = read_book(path, cached=True, checked=True)
    try:
        name = options.get("sheet", "") or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise ValueError("The selected worksheet name does not exist (names are case-sensitive).")
        ws = wb[name]
        rows, columns, characters = 10, 12, 500
        sample = []
        for row in ws.iter_rows(max_row=min(ws.max_row, rows), max_col=min(ws.max_column, columns)):
            cells = []
            for cell in row:
                formula = cell.data_type == "f"
                expression = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)
                fields = {
                    "value": None if formula else plain(cell.value),
                    "formula": expression if formula else None,
                    "cachedValue": plain(cached[name][cell.coordinate].value) if formula else None,
                }
                fields = {key: None if value is None else str(value) for key, value in fields.items()}
                clipped = any(value is not None and len(value) > characters for value in fields.values())
                cells.append({
                    "coordinate": cell.coordinate, "kind": "formula" if formula else "value",
                    **{key: value[:characters] if value is not None else None for key, value in fields.items()},
                    "truncated": clipped,
                })
            sample.append(cells)
        return {
            "sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column,
                        "state": s.sheet_state} for s in wb],
            "previewSheet": name, "sampleRows": sample, "rowLimit": rows,
            "columnLimit": columns, "cellCharacterLimit": characters,
            "truncated": ws.max_row > rows or ws.max_column > columns or any(c["truncated"] for r in sample for c in r),
            "warnings": [
                "The first row is a header candidate, not a detected header. Dimensions include stored formatting.",
                "Formula text is not evaluated. Cached values can be stale; unavailable caches are null.",
                "Only a bounded sample is shown. The original workbook is unchanged.",
            ],
        }
    finally:
        wb.close()
        cached.close()


def saved(wb):
    memory = io.BytesIO()
    wb.save(memory)
    return memory.getvalue()


def simple(wb):
    if wb.defined_names or any(ws.defined_names or ws.tables or ws._charts or ws._images for ws in wb):
        raise ValueError("This operation requires a workbook without named ranges, tables, charts or images.")


def formula(value, mapping, allowed=None):
    if not isinstance(value, str):
        raise ValueError("Array and data-table formulas are unsupported for this operation.")
    try:
        tokens = Tokenizer(value).items
    except Exception:
        raise ValueError("A formula cannot be safely copied.")
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN" and token.value.upper() in {"INDIRECT(", "_XLFN.INDIRECT("}:
            raise ValueError("INDIRECT formulas cannot be safely renamed or split because their references are text.")
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        if "[" in token.value or "]" in token.value:
            raise ValueError("External or structured formula references are unsupported.")
        if "!" in token.value:
            sheet, ref = token.value.rsplit("!", 1)
            original = sheet[1:-1].replace("''", "'") if sheet.startswith("'") and sheet.endswith("'") else sheet
            match = next((name for name in mapping if name.casefold() == original.casefold()), None)
            if not match:
                raise ValueError("A formula references an unknown worksheet or a 3-D sheet range.")
            if allowed is not None and match != allowed:
                raise ValueError("Cannot split a worksheet with cross-sheet formulas. Keep it in the full workbook.")
            token.value = quote_sheetname(mapping[match]) + "!" + ref
        elif not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(:\$?[A-Za-z]{1,3}\$?\d+)?|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d+:\$?\d+", token.value):
            raise ValueError("Named or unsupported formula references cannot be safely copied.")
    return "=" + "".join(t.value for t in tokens)


def copy_sheet(source, dest, mapping, allowed=None):
    for row in source:
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            target = dest.cell(cell.row, cell.column)
            target.value = formula(cell.value, mapping, allowed) if cell.data_type == "f" else cell.value
            target.data_type = cell.data_type
            if cell.has_style:
                target.font, target.fill, target.border = copy(cell.font), copy(cell.fill), copy(cell.border)
                target.alignment, target.protection = copy(cell.alignment), copy(cell.protection)
                target.number_format = cell.number_format
            if cell.comment:
                target.comment = copy(cell.comment)
            if cell.hyperlink:
                target.hyperlink = copy(cell.hyperlink)
    for key, dim in source.row_dimensions.items():
        dest.row_dimensions[key] = copy(dim)
    for key, dim in source.column_dimensions.items():
        dest.column_dimensions[key] = copy(dim)
    for merged in source.merged_cells.ranges:
        dest.merge_cells(str(merged))
    dest.freeze_panes = source.freeze_panes
    dest.auto_filter = copy(source.auto_filter)
    dest.sheet_format = copy(source.sheet_format)
    dest.sheet_properties = copy(source.sheet_properties)
    dest.sheet_view.showGridLines = source.sheet_view.showGridLines
    dest.page_setup, dest.page_margins, dest.print_options = copy(source.page_setup), copy(source.page_margins), copy(source.print_options)
    dest.sheet_state = source.sheet_state


def unique_name(name, used):
    candidate, n = name, 2
    while candidate.casefold() in used:
        suffix = f" ({n})"
        candidate = name[:31-len(suffix)] + suffix
        n += 1
    used.add(candidate.casefold())
    return candidate


def result(files, message, warnings=None, text=""):
    return {"files": files, "message": message, "warnings": warnings or [], "text": text, "pageCount": 0}


def dispatch(slug, paths, options):
    if slug == "inspect-workbook":
        if len(paths) != 1:
            raise ValueError("Choose one XLSX workbook to inspect.")
        return inspect_workbook(paths[0], options)
    response = process_excel(slug, paths, options)
    if response is not None:
        response["originalBytes"] = sum(path.stat().st_size for path in paths)
        response["outputBytes"] = sum(file["bytes"] for file in response["files"])
    return response


def process_excel(slug, paths, options):
    if slug not in SLUGS:
        return None
    if (slug != "merge-excel" and len(paths) != 1) or not paths:
        raise ValueError("Choose one file for this operation.")
    if slug == "csv-to-excel":
        try:
            text = paths[0].read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8 encoded.")
        if "\x00" in text:
            raise ValueError("CSV contains invalid null characters.")
        csv.field_size_limit(32767)
        wb = Workbook()
        count = 0
        try:
            for n, row in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter(options), strict=True), 1):
                count += len(row)
                if n > 100_000 or len(row) > 2000 or count > MAX_CELLS:
                    raise ValueError("CSV exceeds 100,000 rows, 2,000 columns or 500,000 cells.")
                for col, value in enumerate(row, 1):
                    cell = wb.active.cell(n, col, value)
                    cell.data_type = "s"
        except csv.Error:
            raise ValueError("Malformed CSV or a field exceeds 32,767 characters.")
        if not count:
            raise ValueError("CSV contains no fields.")
        return result([output("converted.xlsx", saved(wb))], f"Converted {count:,} fields as literal text.", ["CSV values remain text; formulas, numbers and dates are not inferred."])

    allocation_budget = 0
    for path in paths:
        allocation_budget += package_check(path)
        if allocation_budget > MAX_CELLS:
            raise ValueError("Combined stored cells and merged-cell expansion exceed 500,000 cells.")
    books = []
    dimensions = sheet_count = 0
    for path in paths:
        book = read_book(path, checked=True)
        dimensions += sum(ws.max_row * ws.max_column for ws in book)
        sheet_count += len(book.worksheets)
        if dimensions > MAX_RECTANGLE or sheet_count > MAX_SHEETS:
            raise ValueError("Combined workbooks exceed 500,000 cells or 100 worksheets.")
        books.append(book)
    wb = books[0]
    sheets = selected(wb, options)
    if slug == "merge-excel":
        if len(books) < 2:
            raise ValueError("Choose at least two XLSX workbooks.")
        merged = Workbook()
        merged.remove(merged.active)
        used = set()
        for book in books:
            simple(book)
            mapping = {ws.title: unique_name(ws.title, used) for ws in book}
            for ws in book:
                copy_sheet(ws, merged.create_sheet(mapping[ws.title]), mapping)
        # At least one source workbook already contains a visible sheet.
        return result([output("merged.xlsx", saved(merged))], f"Merged {len(books)} workbooks into {len(merged.worksheets)} worksheets.", WARNINGS + ["Conditional formatting, validation and advanced layout are not copied; duplicate sheet names are renamed."])
    if slug == "split-excel":
        simple(wb)
        parts = []
        mapping = {ws.title: ws.title for ws in wb}
        for i, ws in enumerate(sheets, 1):
            child = Workbook()
            copy_sheet(ws, child.active, mapping, ws.title)
            child.active.title = ws.title
            child.active.sheet_state = "visible"
            parts.append((f"worksheet-{i:03d}.xlsx", saved(child)))
        return result([output("worksheets.zip", zip_outputs(parts))], f"Split {len(parts)} worksheets into workbooks.", WARNINGS + ["Output filenames use worksheet numbers; sheet titles remain in each workbook. Validation and conditional formatting are not copied."])
    if slug == "excel-to-csv":
        mode = options.get("values", "cached")
        if mode not in {"cached", "formula"}:
            raise ValueError("Choose cached values or formula text.")
        cached = read_book(paths[0], cached=True, checked=True)
        parts, missing = [], 0
        for i, ws in enumerate(sheets, 1):
            stream = io.StringIO(newline="")
            writer = csv.writer(stream, delimiter=delimiter(options))
            for row in ws:
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        cache = cached[ws.title][cell.coordinate].value
                        missing += cache is None
                        value = cache if mode == "cached" else (value if isinstance(value, str) else getattr(value, "text", ""))
                    values.append(safe_csv(value))
                writer.writerow(values)
            parts.append((f"worksheet-{i:03d}.csv", b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")))
        return result([output("worksheets-csv.zip", zip_outputs(parts))], f"Exported {len(parts)} worksheets.", [
            f"{missing} formula cells have no cached value; cached mode exports these as blank. Cached values can be stale.",
            "Formula-like CSV text is prefixed with an apostrophe for safety. Formatting is not included.",
        ])
    if slug == "excel-to-json":
        cached = read_book(paths[0], cached=True, checked=True)
        data = {"sheets": []}
        for ws in sheets:
            rows = []
            for row in ws:
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        value = cell.value
                        values.append({"formula": value if isinstance(value, str) else getattr(value, "text", None),
                                       "cachedValue": plain(cached[ws.title][cell.coordinate].value)})
                    else:
                        values.append(plain(cell.value))
                rows.append(values)
            data["sheets"].append({"name": ws.title, "rows": rows})
        encoded = json.dumps(data, ensure_ascii=False, allow_nan=False).encode()
        return result([output("workbook.json", encoded)], "Workbook JSON is ready.", ["Cached formula values can be stale; unavailable caches are null."])
    if slug == "excel-statistics":
        data = {"sheetCount": len(wb.worksheets), "populatedCells": 0, "formulaCells": 0, "sheets": []}
        for ws in wb:
            cells = [c for row in ws for c in row if c.value is not None]
            populated, formulas = len(cells), sum(c.data_type == "f" for c in cells)
            data["populatedCells"] += populated
            data["formulaCells"] += formulas
            data["sheets"].append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                                   "populatedCells": populated, "formulaCells": formulas, "hidden": ws.sheet_state != "visible"})
        text = json.dumps(data, indent=2, ensure_ascii=False)
        return result([output("statistics.json", text.encode())], "Workbook statistics are ready.", ["Dimensions include stored formatting. Formulas were counted, not evaluated."], text)
    simple(wb)
    removed = trimmed = 0
    if slug == "trim-excel-text":
        for ws in sheets:
            for row in ws:
                for cell in row:
                    if cell.data_type != "f" and isinstance(cell.value, str):
                        value = cell.value.strip()
                        if value != cell.value:
                            cell.value = value
                            cell.data_type = "s"
                            trimmed += 1
    else:
        if options.get("header", "yes") not in {"yes", "no"}:
            raise ValueError("Choose whether to retain the header row.")
        if any(c.data_type == "f" for ws in wb for row in ws for c in row):
            raise ValueError("Row removal is unsafe in formula workbooks. Use a values-only copy; the original is unchanged.")
        if any(ws.merged_cells.ranges for ws in wb):
            raise ValueError("Unmerge cells before removing rows.")
        if any(ws.auto_filter.ref or ws.conditional_formatting or ws.data_validations.count for ws in wb):
            raise ValueError("Remove filters, conditional formatting and data validation before deleting rows.")
        if any(ws.print_area or ws.print_title_rows or ws.print_title_cols for ws in wb):
            raise ValueError("Clear print areas and repeating print titles before deleting rows.")
        columns = None
        keep = "first"
        if slug == "remove-excel-duplicates":
            keep = options.get("keep", "first")
            if keep not in {"first", "last"}:
                raise ValueError("Choose first or last for the duplicate row to keep.")
            raw_columns = options.get("columns", "")
            if not isinstance(raw_columns, str):
                raise ValueError("Match columns must be comma-separated column numbers starting at 1.")
            if raw_columns.strip():
                tokens = raw_columns.split(",")
                if len(tokens) > 2000 or any(not re.fullmatch(r"[0-9]{1,4}", token.strip()) for token in tokens):
                    raise ValueError("Match columns must be comma-separated column numbers starting at 1.")
                columns = [int(token.strip()) for token in tokens]
                if len(set(columns)) != len(columns):
                    raise ValueError("Choose each match column only once.")
                for ws in sheets:
                    if any(col < 1 or col > ws.max_column for col in columns):
                        raise ValueError(f"Match column numbers must be between 1 and {ws.max_column} in worksheet '{ws.title}'.")
        for ws in sheets:
            seen, remove = set(), []
            first_row = 2 if options.get("header", "yes") == "yes" else 1
            row_numbers = range(first_row, ws.max_row + 1)
            if keep == "last":
                row_numbers = reversed(row_numbers)
            for n in row_numbers:
                row = ws[n]
                match_cells = (row[col - 1] for col in columns) if columns else row
                key = tuple((c.data_type, type(c.value).__name__, str(plain(c.value))) for c in match_cells)
                empty = all(c.value is None or isinstance(c.value, str) and not c.value.strip() for c in row)
                if (slug == "remove-excel-duplicates" and key in seen) or (slug == "remove-excel-empty-rows" and empty):
                    remove.append(row[0].row)
                seen.add(key)
            remove.sort()
            if remove:
                removed_set = set(remove)
                cells = {}
                for (n, col), cell in ws._cells.items():
                    if n not in removed_set:
                        cell.row = n - bisect_left(remove, n)
                        if cell.hyperlink:
                            cell.hyperlink.ref = cell.coordinate
                        cells[(cell.row, col)] = cell
                ws._cells = cells
                ws._current_row = ws.max_row
                if ws.freeze_panes:
                    n, col = coordinate_to_tuple(ws.freeze_panes)
                    ws.freeze_panes = ws.cell(max(1, n - bisect_left(remove, n)), col).coordinate
            removed += len(remove)
            # Row dimensions are not shifted by openpyxl.delete_rows.
            if remove:
                dimensions = {n: copy(d) for n, d in ws.row_dimensions.items()}
                ws.row_dimensions.clear()
                for n, dim in dimensions.items():
                    if n not in removed_set:
                        new = n - bisect_left(remove, n)
                        dim.index = new
                        ws.row_dimensions[new] = dim
    return result([output("cleaned.xlsx", saved(wb))], f"Trimmed {trimmed} text cells." if slug == "trim-excel-text" else f"Removed {removed} rows.", WARNINGS)