"""Reopen real results, including hostile inputs and formula safety."""
import base64
import csv
import io
import json
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from pathlib import Path
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_tools import dispatch


class ExcelToolsTests(unittest.TestCase):
    def test_inspection_bounds_and_sheet_picker_metadata(self):
        wb, path = self.book(rows=[["x" * 900] * 15 for _ in range(25)])
        hidden = wb.create_sheet("Sales & 日本語")
        hidden.sheet_state = "hidden"
        hidden.append(["Header", "=1+1", "=literal", date(2026, 1, 2)])
        hidden["C1"].data_type = "s"
        wb.save(path)
        original = path.read_bytes()
        preview = dispatch("inspect-workbook", [path], {})
        self.assertEqual(preview["previewSheet"], "Data")
        self.assertEqual(len(preview["sampleRows"]), 10)
        self.assertTrue(all(len(row) == 12 for row in preview["sampleRows"]))
        self.assertTrue(preview["truncated"])
        self.assertEqual(len(preview["sampleRows"][0][0]["value"]), 500)
        self.assertTrue(preview["sampleRows"][0][0]["truncated"])
        self.assertEqual(preview["sheets"][1], {
            "name": "Sales & 日本語", "rows": 1, "columns": 4, "state": "hidden",
        })
        selected = dispatch("inspect-workbook", [path], {"sheet": "Sales & 日本語"})
        cells = selected["sampleRows"][0]
        self.assertEqual(cells[1]["kind"], "formula")
        self.assertEqual(cells[1]["formula"], "=1+1")
        self.assertIsNone(cells[1]["value"])
        self.assertIsNone(cells[1]["cachedValue"])
        self.assertEqual(cells[2]["kind"], "value")
        self.assertEqual(cells[2]["value"], "=literal")
        self.assertEqual(cells[3]["value"], "2026-01-02T00:00:00")
        self.assertEqual(path.read_bytes(), original)
        json.dumps(selected, allow_nan=False)

    def test_inspection_stored_cache_is_not_recalculated(self):
        _, path = self.book(rows=[["Result"], ["=1+1"]])
        with zipfile.ZipFile(path) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        # Deliberately stale cache: inspection must show 99, not evaluate to 2.
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
            b"<f>1+1</f><v></v>", b"<f>1+1</f><v>99</v>")
        with zipfile.ZipFile(path, "w") as target:
            for name, content in parts.items():
                target.writestr(name, content)
        cell = dispatch("inspect-workbook", [path], {})["sampleRows"][1][0]
        self.assertEqual(cell["formula"], "=1+1")
        self.assertEqual(cell["cachedValue"], "99")

    def test_inspection_invalid_sheet_and_input(self):
        _, path = self.book()
        with self.assertRaisesRegex(ValueError, "does not exist"):
            dispatch("inspect-workbook", [path], {"sheet": "data"})
        with self.assertRaisesRegex(ValueError, "one XLSX"):
            dispatch("inspect-workbook", [path, path], {})
        invalid = self.root / "invalid.xlsx"
        invalid.write_bytes(b"not an xlsx")
        with self.assertRaisesRegex(ValueError, "Malformed"):
            dispatch("inspect-workbook", [invalid], {})
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"unsupported macro")
        with patch("excel_tools.load_workbook") as loader:
            with self.assertRaisesRegex(ValueError, "Macros"):
                dispatch("inspect-workbook", [path], {})
            loader.assert_not_called()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def book(self, name="test.xlsx", rows=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        for row in rows or [["name", "value"], ["Alpha", 2], ["Beta", 3]]:
            ws.append(row)
        path = self.root / name
        wb.save(path)
        return wb, path

    def run_tool(self, slug, paths, **options):
        result = dispatch(slug, paths, options)
        self.assertEqual(result["originalBytes"], sum(path.stat().st_size for path in paths))
        self.assertEqual(result["outputBytes"], sum(file["bytes"] for file in result["files"]))
        return result

    def data(self, result):
        file = result["files"][0]
        decoded = base64.b64decode(file["dataBase64"])
        self.assertEqual(file["bytes"], len(decoded))
        return decoded

    def reopen(self, result):
        return load_workbook(io.BytesIO(self.data(result)))

    def test_csv_literal_text(self):
        path = self.root / "text.csv"
        path.write_text('id,value\n001,=1+1\n002,+cmd\n003,@SUM(A1)\n004,-2\n005,"hello, there"\n', encoding="utf-8")
        wb = self.reopen(self.run_tool("csv-to-excel", [path]))
        self.assertEqual(wb.active["A2"].value, "001")
        for row in range(2, 6):
            self.assertEqual(wb.active[f"B{row}"].data_type, "s")
        self.assertEqual(wb.active["B2"].value, "=1+1")
        self.assertEqual(wb.active["B6"].value, "hello, there")

    def test_csv_delimiters_and_bom(self):
        path = self.root / "text.csv"
        path.write_bytes(b"\xef\xbb\xbfid;value\r\n01;test\r\n")
        wb = self.reopen(self.run_tool("csv-to-excel", [path], delimiter="semicolon"))
        self.assertEqual(wb.active["B2"].value, "test")

    def test_merge_renames_references_preserves_styles(self):
        wb, one = self.book("one.xlsx")
        wb.active["C2"] = "=SUM(B2:B3)"
        wb.active["A2"].font = Font(bold=True)
        wb.active.column_dimensions["A"].width = 24
        wb.active.freeze_panes = "A2"
        wb.save(one)
        other, two = self.book("two.xlsx")
        other.active["C2"] = "='Data'!B2+Data!B3"
        other.save(two)
        merged = self.reopen(self.run_tool("merge-excel", [one, two]))
        self.assertEqual(merged.sheetnames, ["Data", "Data (2)"])
        self.assertEqual(merged["Data"]["C2"].value, "=SUM(B2:B3)")
        self.assertEqual(merged["Data (2)"]["C2"].value, "='Data (2)'!B2+'Data (2)'!B3")
        self.assertTrue(merged["Data"]["A2"].font.bold)
        self.assertEqual(merged["Data"].column_dimensions["A"].width, 24)
        self.assertEqual(merged["Data"].freeze_panes, "A2")

    def test_split_reopens_and_formula(self):
        wb, path = self.book()
        wb.active["C2"] = "=B2*2"
        wb.create_sheet("Extra").append(["hello"])
        wb.save(path)
        with zipfile.ZipFile(io.BytesIO(self.data(self.run_tool("split-excel", [path])))) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            first = load_workbook(io.BytesIO(archive.read("worksheet-001.xlsx")))
            self.assertEqual(first.active.title, "Data")
            self.assertEqual(first.active["C2"].value, "=B2*2")

    def test_split_cross_sheet_rejected(self):
        wb, path = self.book()
        wb.create_sheet("Extra")["A1"] = "=Data!B2"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "cross-sheet"):
            self.run_tool("split-excel", [path])

    def test_trim_preserves_formula_and_literal_equal(self):
        wb, path = self.book()
        wb.active["A2"] = "  =1+1  "
        wb.active["C2"] = "=B2*2"
        wb.save(path)
        trimmed = self.reopen(self.run_tool("trim-excel-text", [path]))
        self.assertEqual(trimmed.active["A2"].value, "=1+1")
        self.assertEqual(trimmed.active["A2"].data_type, "s")
        self.assertEqual(trimmed.active["C2"].data_type, "f")
        self.assertEqual(trimmed.active["B2"].value, 2)

    def test_duplicate_and_empty_rows(self):
        wb, path = self.book(rows=[["name", "value"], ["x", 2], ["x", 2], [None, " "], ["zero", 0], ["boolean", False]])
        wb.active["A5"].font = Font(bold=True)
        wb.active.row_dimensions[5].height = 30
        wb.save(path)
        clean = self.reopen(self.run_tool("remove-excel-duplicates", [path]))
        self.assertEqual(clean.active.max_row, 5)
        self.assertEqual(clean.active["A4"].value, "zero")
        self.assertTrue(clean.active["A4"].font.bold)
        self.assertEqual(clean.active.row_dimensions[4].height, 30)
        clean = self.reopen(self.run_tool("remove-excel-empty-rows", [path]))
        self.assertEqual(clean.active.max_row, 5)
        self.assertEqual(clean.active["B4"].value, 0)
        self.assertIs(clean.active["B5"].value, False)

    def test_header_and_sheet_selection(self):
        _, path = self.book(rows=[["same"], ["same"], ["same"]])
        clean = self.reopen(self.run_tool("remove-excel-duplicates", [path], header="no", sheet="Data"))
        self.assertEqual(clean.active.max_row, 1)
        with self.assertRaisesRegex(ValueError, "does not exist"):
            self.run_tool("excel-to-json", [path], sheet="Missing")

    def test_selected_column_duplicates_keep_first_and_last(self):
        rows = [["ID", "Name"], ["a", "old"], ["b", "only"], ["a", "new"], ["ID", "customer"]]
        wb, path = self.book(rows=rows)
        wb.active["B4"].font = Font(bold=True)
        wb.active["B4"].hyperlink = "https://example.com"
        wb.active.row_dimensions[4].height = 31
        wb.active.freeze_panes = "B4"
        wb.save(path)
        original = path.read_bytes()
        for keep, expected in (
            ("first", [rows[0], rows[1], rows[2], rows[4]]),
            ("last", [rows[0], rows[2], rows[3], rows[4]]),
        ):
            with self.subTest(keep=keep):
                result = self.run_tool("remove-excel-duplicates", [path], columns="1", keep=keep)
                ws = self.reopen(result).active
                self.assertEqual(list(ws.values), [tuple(r) for r in expected])
                self.assertEqual(result["message"], "Removed 1 rows.")
                if keep == "last":
                    self.assertTrue(ws["B3"].font.bold)
                    self.assertEqual(ws.row_dimensions[3].height, 31)
                    self.assertEqual(ws["B3"].hyperlink.ref, "B3")
                    self.assertEqual(ws.freeze_panes, "B3")
        self.assertEqual(path.read_bytes(), original)
        # Default matching still compares all cells, not just IDs.
        self.assertEqual(list(self.reopen(self.run_tool("remove-excel-duplicates", [path])).active.values),
                         [tuple(r) for r in rows])

    def test_duplicate_composite_keys_and_header_choices(self):
        rows = [["same", "ignored", "key"], ["same", "first", "key"],
                ["same", "different", "other"], ["same", "last", "key"]]
        _, path = self.book(rows=rows)
        for header, keep, indexes in (
            ("yes", "first", [0, 1, 2]), ("yes", "last", [0, 2, 3]),
            ("no", "first", [0, 2]), ("no", "last", [2, 3]),
        ):
            with self.subTest(header=header, keep=keep):
                clean = self.reopen(self.run_tool("remove-excel-duplicates", [path],
                                                 columns=" 3, 1 ", header=header, keep=keep))
                self.assertEqual(list(clean.active.values), [tuple(rows[i]) for i in indexes])

    def test_duplicate_keys_distinguish_types_case_whitespace_and_empty_values(self):
        keys = [1, "1", True, 0, False, None, "", " ", "a", "A", " a", date(2026, 1, 2), "2026-01-02"]
        rows = [["Key", "Marker"]]
        for i, key in enumerate(keys):
            rows.extend([[key, f"first-{i}"], [key, f"last-{i}"]])
        _, path = self.book(rows=rows)
        for keep in ("first", "last"):
            ws = self.reopen(self.run_tool("remove-excel-duplicates", [path], columns="1", keep=keep)).active
            self.assertEqual(ws.max_row, len(keys) + 1)
            self.assertEqual([r[1] for r in list(ws.values)[1:]], [f"{keep}-{i}" for i in range(len(keys))])
            self.assertEqual(ws["A2"].data_type, "n")
            self.assertEqual(ws["A3"].data_type, "s")
            self.assertEqual(ws["A4"].data_type, "b")

    def test_duplicate_matching_is_per_sheet_and_validates_all_selected_sheets(self):
        wb, path = self.book(rows=[["ID", "Name"], ["a", "old"], ["a", "new"]])
        other = wb.create_sheet("Other")
        for row in [["ID"], ["a"], ["a"]]:
            other.append(row)
        wb.save(path)
        for sheet, remaining in (("", 2), ("Data", 3)):
            clean = self.reopen(self.run_tool("remove-excel-duplicates", [path], columns="1", keep="last", sheet=sheet))
            self.assertEqual(clean["Data"]["B2"].value, "new")
            self.assertEqual(clean["Other"].max_row, remaining)
        original = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "worksheet 'Other'"):
            self.run_tool("remove-excel-duplicates", [path], columns="2")
        self.assertEqual(path.read_bytes(), original)
        self.reopen(self.run_tool("remove-excel-duplicates", [path], columns="2", sheet="Data"))

    def test_duplicate_matching_rejects_invalid_options(self):
        _, path = self.book()
        for columns in ("0", "-1", "3", "1.5", "A", "1,", ",1", "1 2", "1,1", "1,01",
                        "999999999999999", "1;2", "1," * 2001, ["1"], None):
            with self.subTest(columns=columns), self.assertRaisesRegex(ValueError, "[Cc]olumn"):
                self.run_tool("remove-excel-duplicates", [path], columns=columns)
        with self.assertRaisesRegex(ValueError, "first or last"):
            self.run_tool("remove-excel-duplicates", [path], keep="newest")
        for columns in ("", "   "):
            self.assertEqual(self.reopen(self.run_tool("remove-excel-duplicates", [path], columns=columns)).active.max_row, 3)

    def test_duplicate_all_cell_last_and_header_only(self):
        _, path = self.book(rows=[["header"], ["a"], ["b"], ["a"]])
        clean = self.reopen(self.run_tool("remove-excel-duplicates", [path], keep="last"))
        self.assertEqual(list(clean.active.values), [("header",), ("b",), ("a",)])
        for rows in ([["header"]], [[None, None]]):
            _, path = self.book(rows=rows)
            for keep in ("first", "last"):
                clean = self.reopen(self.run_tool("remove-excel-duplicates", [path], columns="1", keep=keep))
                self.assertEqual(clean.active.max_row, 1)

    def test_selected_column_duplicates_reject_formulas_even_on_unselected_sheets(self):
        wb, path = self.book(rows=[["ID", "Value"], ["a", 1], ["a", 2]])
        wb.create_sheet("Formula")["A1"] = "=Data!B2"
        wb.save(path)
        original = path.read_bytes()
        for keep in ("first", "last"):
            with self.assertRaisesRegex(ValueError, "formula workbooks"):
                self.run_tool("remove-excel-duplicates", [path], sheet="Data", columns="1", keep=keep)
        self.assertEqual(path.read_bytes(), original)

    def test_row_cleanup_rejects_formulas(self):
        wb, path = self.book()
        wb.active["C2"] = "=B2"
        wb.save(path)
        for slug in ("remove-excel-duplicates", "remove-excel-empty-rows"):
            with self.assertRaisesRegex(ValueError, "formula workbooks"):
                self.run_tool(slug, [path])

    def test_json_distinguishes_formula_cache_and_dates(self):
        wb, path = self.book()
        wb.active["C2"] = "=B2*2"
        wb.active["D2"] = date(2026, 10, 4)
        wb.save(path)
        data = json.loads(self.data(self.run_tool("excel-to-json", [path])))
        values = data["sheets"][0]["rows"][1]
        self.assertEqual(values[2], {"formula": "=B2*2", "cachedValue": None})
        self.assertTrue(values[3].startswith("2026-10-04"))
        # Add an actual stored cache without evaluating anything.
        self.rewrite(path, lambda name, data: data.replace(b"<v></v>", b"<v>4</v>", 1) if name == "xl/worksheets/sheet1.xml" else data)
        values = json.loads(self.data(self.run_tool("excel-to-json", [path])))["sheets"][0]["rows"][1]
        self.assertEqual(values[2]["cachedValue"], 4)

    def test_csv_export_safe_formula_and_cache(self):
        wb, path = self.book()
        wb.active["A2"] = "\t=cmd"
        wb.active["C2"] = "=B2*2"
        wb.save(path)
        for mode in ("cached", "formula"):
            result = self.run_tool("excel-to-csv", [path], values=mode)
            with zipfile.ZipFile(io.BytesIO(self.data(result))) as archive:
                rows = list(csv.reader(io.StringIO(archive.read("worksheet-001.csv").decode("utf-8-sig"))))
            self.assertTrue(rows[1][0].startswith("'"))
            self.assertEqual(rows[1][2], "'=B2*2" if mode == "formula" else "")
            self.assertIn("1 formula", result["warnings"][0])

    def test_statistics_real_counts(self):
        wb, path = self.book()
        wb.active["C2"] = "=B2*2"
        wb.save(path)
        result = self.run_tool("excel-statistics", [path])
        stats = json.loads(result["text"])
        self.assertEqual(stats["formulaCells"], 1)
        self.assertEqual(stats["populatedCells"], 7)
        self.assertEqual(json.loads(self.data(result)), stats)

    def rewrite(self, path, transform):
        with zipfile.ZipFile(path) as source:
            entries = [(name, transform(name, source.read(name))) for name in source.namelist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
            for name, data in entries:
                target.writestr(name, data)

    def test_malformed_and_hostile_inputs(self):
        path = self.root / "bad.xlsx"
        path.write_bytes(b"not a workbook")
        with self.assertRaises(ValueError):
            self.run_tool("excel-statistics", [path])
        for part in ("../escape.xml", "xl/vbaProject.bin", "xl/externalLinks/link.xml"):
            _, path = self.book()
            with zipfile.ZipFile(path, "a") as archive:
                archive.writestr(part, b"<root/>")
            with self.assertRaises(ValueError):
                self.run_tool("excel-statistics", [path])
        _, path = self.book()
        self.rewrite(path, lambda name, data: b'<!DOCTYPE x [<!ENTITY y "evil">]><x/>' if name == "xl/workbook.xml" else data)
        with self.assertRaisesRegex(ValueError, "XML declarations"):
            self.run_tool("excel-statistics", [path])
        csv_path = self.root / "bad.csv"
        for content in (b"\xff\xfe", b"a,\x00", b'a,"unclosed'):
            csv_path.write_bytes(content)
            with self.assertRaises(ValueError):
                self.run_tool("csv-to-excel", [csv_path])

    def test_oversized_dimensions_and_invalid_options(self):
        wb, path = self.book()
        wb.active["XFD1048576"] = "distant"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "exceeds"):
            self.run_tool("excel-statistics", [path])
        _, path = self.book()
        with self.assertRaises(ValueError):
            self.run_tool("merge-excel", [path])
        with self.assertRaises(ValueError):
            self.run_tool("excel-to-csv", [path], delimiter="bad")
        with self.assertRaises(ValueError):
            self.run_tool("excel-to-csv", [path], values="bad")

    def test_complex_workbooks_rejected_not_silently_lost(self):
        from openpyxl.worksheet.table import Table
        wb, path = self.book()
        wb.active.add_table(Table(displayName="DataTable", ref="A1:B3"))
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "tables"):
            self.run_tool("split-excel", [path])
        wb, path = self.book()
        wb.active.merge_cells("A1:B1")
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "Unmerge"):
            self.run_tool("remove-excel-empty-rows", [path])

    def test_indirect_and_print_references_rejected(self):
        wb, path = self.book()
        wb.active["C2"] = '=INDIRECT("Data!B2")'
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "INDIRECT"):
            self.run_tool("split-excel", [path])
        wb, path = self.book()
        wb.active.print_area = "A1:B3"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "print areas"):
            self.run_tool("remove-excel-empty-rows", [path])

    def test_hostile_merged_ranges_rejected_before_loading(self):
        # Inject XML directly; constructing these ranges in openpyxl would
        # already allocate the cells we are specifically preventing.
        ranges = [
            ["A1:XFD1048576"],
            ["A1:ALL100000"],  # Within per-axis limits but huge area.
            ["A1:Z10000", "AA1:AZ10000"],  # Cumulative expansion.
            ["A1:Z10000", "A1:Z10000"],  # Overlap still costs allocations.
            ["A:A"], ["1:3"], ["B2:A1"], ["A0:B1"], ["not-a-range"],
        ]
        for refs in ranges:
            with self.subTest(refs=refs):
                _, path = self.book()
                xml = ('<mergeCells count="%d">' % len(refs) +
                       "".join(f'<mergeCell ref="{ref}"/>' for ref in refs) + "</mergeCells>").encode()
                self.rewrite(path, lambda name, data: data.replace(b"</worksheet>", xml + b"</worksheet>") if name == "xl/worksheets/sheet1.xml" else data)
                with patch("excel_tools.load_workbook") as loader:
                    with self.assertRaisesRegex(ValueError, "Merged-cell|merged-cell"):
                        self.run_tool("excel-statistics", [path])
                    loader.assert_not_called()

    def test_small_merged_range_still_loads_and_roundtrips(self):
        wb, path = self.book()
        wb.active.merge_cells("A1:B1")
        wb.save(path)
        result = self.run_tool("trim-excel-text", [path])
        reopened = self.reopen(result)
        self.assertEqual(str(next(iter(reopened.active.merged_cells.ranges))), "A1:B1")

    def relocate_sheet(self, path, sheet_xml):
        with zipfile.ZipFile(path) as source:
            parts = []
            for name in source.namelist():
                data = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    name, data = "xl/custom/sheet1.xml", sheet_xml
                elif name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                    data = data.replace(b"/xl/worksheets/sheet1.xml", b"/xl/custom/sheet1.xml")
                    data = data.replace(b'Target="worksheets/sheet1.xml"', b'Target="custom/sheet1.xml"')
                parts.append((name, data))
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
            for name, data in parts:
                target.writestr(name, data)

    def test_relocated_worksheet_rejected_before_loader(self):
        for content in (
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/><mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>',
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/><mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells></worksheet>',
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">' + b'<c r="A1"><v>1</v></c>' * 500_001 + b'</row></sheetData></worksheet>',
        ):
            _, path = self.book()
            self.relocate_sheet(path, content)
            with patch("excel_tools.load_workbook") as loader:
                with self.assertRaisesRegex(ValueError, "worksheet relationship location"):
                    self.run_tool("excel-statistics", [path])
                loader.assert_not_called()

    def test_aggregate_expansion_budget_rejected_before_loader(self):
        paths = []
        for i in range(2):
            _, path = self.book(f"book{i}.xlsx")
            self.rewrite(path, lambda name, data: data.replace(b"</worksheet>", b'<mergeCells count="1"><mergeCell ref="A1:Z10000"/></mergeCells></worksheet>') if name == "xl/worksheets/sheet1.xml" else data)
            paths.append(path)
        with patch("excel_tools.load_workbook") as loader:
            with self.assertRaisesRegex(ValueError, "Combined stored cells"):
                self.run_tool("merge-excel", paths)
            loader.assert_not_called()

    def test_manifest_cannot_redirect_workbook(self):
        _, path = self.book()
        with zipfile.ZipFile(path) as source:
            relocated = source.read("xl/workbook.xml")
            rels = source.read("xl/_rels/workbook.xml.rels")
        self.rewrite(path, lambda name, data: data.replace(b"/xl/workbook.xml", b"/xl/custom/book.xml") if name == "[Content_Types].xml" else data)
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/custom/book.xml", relocated)
            archive.writestr("xl/custom/_rels/book.xml.rels", rels)
        with patch("excel_tools.load_workbook") as loader:
            with self.assertRaisesRegex(ValueError, "standard XLSX workbook package locations"):
                self.run_tool("excel-statistics", [path])
            loader.assert_not_called()

    def test_repeated_worksheet_target_rejected_before_loader(self):
        wb, path = self.book()
        wb.create_sheet("Second")["A1"] = "hello"
        wb.save(path)
        self.rewrite(path, lambda name, data: data.replace(b"worksheets/sheet2.xml", b"worksheets/sheet1.xml") if name == "xl/_rels/workbook.xml.rels" else data)
        with patch("excel_tools.load_workbook") as loader:
            with self.assertRaisesRegex(ValueError, "same worksheet part"):
                self.run_tool("excel-statistics", [path])
            loader.assert_not_called()


if __name__ == "__main__":
    unittest.main()