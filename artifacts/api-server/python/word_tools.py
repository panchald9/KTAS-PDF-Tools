"""Local DOCX operations. Never extract packages or fetch external resources."""
import csv
import io
import json
import math
import re
import zipfile
from html import escape
from pathlib import PurePosixPath

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree

from common import output

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
WORD_SLUGS = {
    "text-to-word", "merge-word", "word-to-text", "word-to-html",
    "word-tables-to-csv", "extract-word-images", "word-statistics", "word-metadata",
    "remove-word-metadata", "remove-word-comments", "accept-word-revisions",
    "find-replace-word", "word-page-size", "word-margins", "word-font",
    "word-line-spacing", "word-header-footer", "remove-word-empty-paragraphs",
}


def read_package(path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 4000 or sum(e.file_size for e in entries) > 120 * 1024 * 1024:
                raise ValueError("DOCX package is too large when expanded.")
            names = [e.filename for e in entries]
            if len(set(names)) != len(names):
                raise ValueError("DOCX contains duplicate package parts.")
            for item in entries:
                name = PurePosixPath(item.filename)
                if name.is_absolute() or ".." in name.parts or "\\" in item.filename or item.flag_bits & 1:
                    raise ValueError("Invalid or encrypted DOCX package.")
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ValueError("This file is not a Word DOCX document.")
            package = {e.filename: archive.read(e) for e in entries}
    except (zipfile.BadZipFile, KeyError):
        raise ValueError("The DOCX package cannot be read.")
    for name, data in package.items():
        if name.endswith((".xml", ".rels")):
            if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                raise ValueError("DOCX contains unsupported XML declarations.")
            try:
                parse(data)
            except etree.XMLSyntaxError:
                raise ValueError("DOCX contains malformed XML.")
    return package


def parse(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def xml_bytes(root):
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def packed(parts):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return stream.getvalue()


def document_bytes(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def bound(options, key, default, low, high):
    try:
        value = float(options.get(key, str(default)))
    except (ValueError, TypeError):
        raise ValueError(f"{key} must be a number.")
    if not math.isfinite(value) or not low <= value <= high:
        raise ValueError(f"{key} must be between {low} and {high}.")
    return value


def body_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from body_paragraphs(cell)


def all_paragraphs(doc):
    yield from body_paragraphs(doc)
    seen = set()
    for section in doc.sections:
        for region in (section.header, section.footer, section.first_page_header,
                       section.first_page_footer, section.even_page_header, section.even_page_footer):
            # Do not create previously absent header/footer parts just by inspecting them.
            if region._has_definition and region.part.partname not in seen:
                seen.add(region.part.partname)
                yield from body_paragraphs(region)


def body_text(doc):
    blocks = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            blocks.append("".join(child.xpath(".//w:t/text()")))
        elif child.tag == qn("w:tbl"):
            for row in child.findall(qn("w:tr")):
                blocks.append("\t".join(" ".join("".join(p.xpath(".//w:t/text()"))
                                               for p in cell.findall(qn("w:p")))
                                         for cell in row.findall(qn("w:tc"))))
    return "\n".join(blocks)


def replace_paragraph(p, find, replacement):
    # Keep the original runs, drawings, hyperlinks and formatting; change only text nodes.
    nodes, chunks, cursor = [], [], 0
    for node in p.iter():
        nearest = next((a for a in node.iterancestors() if a.tag == qn("w:p")), None)
        if nearest is not p or any(a.tag == qn("w:txbxContent") for a in node.iterancestors()):
            continue
        if node.tag == qn("w:t"):
            value = node.text or ""
            nodes.append((node, cursor, cursor + len(value)))
        elif node.tag == qn("w:tab"):
            value = "\t"
        elif node.tag in (qn("w:br"), qn("w:cr")):
            value = "\n"
        else:
            continue
        chunks.append(value)
        cursor += len(value)
    text = "".join(chunks)
    matches = list(re.finditer(re.escape(find), text))
    changed = 0
    for match in reversed(matches):
        selected = [(n, start, end) for n, start, end in nodes if start < match.end() and end > match.start()]
        if not selected:
            continue
        # Tabs and breaks are not editable text nodes.
        if sum(min(end, match.end()) - max(start, match.start()) for _, start, end in selected) != len(find):
            continue
        for index, (node, start, end) in enumerate(selected):
            value = node.text or ""
            left = max(0, match.start() - start)
            right = min(end, match.end()) - start
            node.text = value[:left] + (replacement if index == 0 else "") + value[right:]
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        changed += 1
    return changed


def raw_edit(parts, slug, options):
    count = 0
    if slug == "remove-word-comments":
        removed = {n for n in parts if PurePosixPath(n).name.lower().startswith(
            ("comments", "people."))}
        for name in removed:
            del parts[name]
        for name, data in list(parts.items()):
            if not name.endswith((".xml", ".rels")):
                continue
            root = parse(data)
            for node in list(root.iter()):
                local = etree.QName(node).localname
                target = node.get("Target", "").lower()
                reltype = node.get("Type", "").lower()
                partname = node.get("PartName", "").lstrip("/")
                remove = local in {"commentRangeStart", "commentRangeEnd", "commentReference"}
                remove |= name.endswith(".rels") and ("comment" in reltype or reltype.endswith("/people")
                                                       or PurePosixPath(target).name.startswith(("comments", "people.")))
                remove |= name == "[Content_Types].xml" and partname in removed
                if remove and node.getparent() is not None:
                    node.getparent().remove(node)
                    count += 1
            parts[name] = xml_bytes(root)
    else:
        find = options.get("find", "")
        if slug == "find-replace-word" and (not find or len(find) > 10000):
            raise ValueError("Enter text to find (up to 10,000 characters).")
        for name, data in list(parts.items()):
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = parse(data)
            if slug == "find-replace-word":
                if not (name == "word/document.xml" or re.match(r"word/(header|footer)\d+\.xml$", name)):
                    continue
                for p in root.findall(".//w:p", NS):
                    count += replace_paragraph(p, find, options.get("replace", ""))
            elif slug == "accept-word-revisions":
                for node in list(root.iter()):
                    if node.getparent() is None or etree.QName(node).namespace != NS["w"]:
                        continue
                    local = etree.QName(node).localname
                    parent = node.getparent()
                    if local in {"del", "moveFrom"}:
                        # Row/cell deletion markers belong to their enclosing structural element.
                        structural = parent.getparent() if parent.tag in {qn("w:trPr"), qn("w:tcPr")} else None
                        if structural is not None and structural.getparent() is not None:
                            structural.getparent().remove(structural)
                        else:
                            parent.remove(node)
                        count += 1
                    elif local in {"ins", "moveTo"}:
                        index = parent.index(node)
                        for child in list(node):
                            parent.insert(index, child)
                            index += 1
                        parent.remove(node)
                        count += 1
                    elif local.endswith("PrChange") or local in {
                        "moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd",
                        "cellIns", "cellMerge", "numberingChange",
                    }:
                        parent.remove(node)
                        count += 1
                    elif local == "cellDel":
                        cell = parent.getparent()
                        if cell is not None and cell.getparent() is not None:
                            cell.getparent().remove(cell)
                        count += 1
            parts[name] = xml_bytes(root)
    return packed(parts), count


def dispatch(slug, paths, options):
    if slug not in WORD_SLUGS:
        return None
    if slug == "merge-word" and len(paths) < 2:
        raise ValueError("Upload at least two DOCX files to merge.")
    if slug != "merge-word" and len(paths) != 1:
        raise ValueError("Upload exactly one file.")
    original = sum(p.stat().st_size for p in paths)
    files, text, warnings = [], "", []
    message = "Your Word document has been processed. Open and check the result."
    data, extension = None, ".docx"
    packages = [] if slug == "text-to-word" else [read_package(p) for p in paths]
    # Prevent a multi-document request expanding beyond the worker memory budget.
    if sum(len(value) for package in packages for value in package.values()) > 150 * 1024 * 1024:
        raise ValueError("Combined expanded DOCX packages are too large.")
    if slug == "text-to-word":
        try:
            content = paths[0].read_text(encoding="utf-8-sig")
        except UnicodeError:
            raise ValueError("Upload a UTF-8 encoded text file.")
        if len(content) > 1_000_000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", content):
            raise ValueError("Text must be at most one million characters and contain no unsupported control characters.")
        doc = Document()
        if options.get("title"):
            doc.add_heading(options["title"], 0)
        for line in content.splitlines():
            doc.add_paragraph(line)
        data = document_bytes(doc)
    elif slug == "merge-word":
        from docxcompose.composer import Composer
        composer = Composer(Document(io.BytesIO(packed(packages[0]))))
        for package in packages[1:]:
            composer.doc.add_page_break()
            composer.append(Document(io.BytesIO(packed(package))))
        stream = io.BytesIO()
        composer.save(stream)
        data = stream.getvalue()
        warnings.append("Headers and footers are taken from the first document. Check sections, fields, numbering and styles.")
    elif slug in {"remove-word-comments", "accept-word-revisions", "find-replace-word"}:
        data, count = raw_edit(packages[0], slug, options)
        message = f"Processed {count} matching items. Download and review the document."
    else:
        parts = packages[0]
        doc = Document(io.BytesIO(packed(parts)))
        if slug == "word-to-text":
            text = body_text(doc)
            data, extension = (text or "\n").encode("utf-8"), ".txt"
        elif slug == "word-to-html":
            blocks = []
            for child in doc.element.body:
                if child.tag == qn("w:p"):
                    blocks.append("<p>" + escape("".join(child.xpath(".//w:t/text()"))) + "</p>")
                elif child.tag == qn("w:tbl"):
                    blocks.append("<table>")
                    for row in child.findall(qn("w:tr")):
                        blocks.append("<tr>" + "".join("<td>" + escape(" ".join(cell.xpath(".//w:t/text()"))) + "</td>"
                                                      for cell in row.findall(qn("w:tc"))) + "</tr>")
                    blocks.append("</table>")
            data = ('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
                    '<title>Word document export</title><style>body{max-width:900px;margin:40px auto;padding:20px;font-family:system-ui}'
                    'p{white-space:pre-wrap}table{border-collapse:collapse}td{border:1px solid #888;padding:8px}</style><body>'
                    + "".join(blocks) + "</body></html>").encode("utf-8")
            extension = ".html"
        elif slug == "word-tables-to-csv":
            if not doc.tables:
                raise ValueError("No top-level tables were found in this document.")
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                for i, table in enumerate(doc.tables):
                    csvtext = io.StringIO(newline="")
                    writer = csv.writer(csvtext)
                    for row in table.rows:
                        values = [cell.text for cell in row.cells]
                        writer.writerow(["'" + value if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value
                                         for value in values])
                    archive.writestr(f"table-{i+1}.csv", csvtext.getvalue().encode("utf-8-sig"))
            data, extension = buffer.getvalue(), ".zip"
        elif slug == "extract-word-images":
            images = [(name, value) for name, value in parts.items() if name.startswith("word/media/")
                      and PurePosixPath(name).suffix.lower() in {
                          ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".svg", ".emf", ".wmf", ".webp"}]
            if not images:
                raise ValueError("No embedded images were found.")
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                for i, (name, value) in enumerate(images):
                    archive.writestr(f"image-{i+1}{PurePosixPath(name).suffix.lower()}", value)
            data, extension = buffer.getvalue(), ".zip"
        elif slug in {"word-statistics", "word-metadata"}:
            if slug == "word-statistics":
                content = body_text(doc)
                info = {"words": len(content.split()), "characters": len(content),
                        "charactersWithoutWhitespace": len(re.sub(r"\s", "", content)),
                        "bodyParagraphs": len(doc.paragraphs), "topLevelTables": len(doc.tables),
                        "embeddedImages": len([n for n in parts if n.startswith("word/media/")]),
                        "pageCount": None, "note": "Page count requires rendering; no estimate is made."}
            else:
                core = doc.core_properties
                info = {key: str(getattr(core, key) or "") for key in (
                    "title", "subject", "author", "keywords", "comments", "last_modified_by",
                    "category", "content_status", "created", "modified", "revision", "language", "version", "identifier")}
            text = json.dumps(info, ensure_ascii=False, indent=2)
            data, extension = text.encode("utf-8"), ".json"
        elif slug == "remove-word-metadata":
            # Raw core-part edits avoid setting artificial created/modified dates on save.
            if "docProps/core.xml" in parts:
                root = parse(parts["docProps/core.xml"])
                for child in list(root):
                    root.remove(child)
                parts["docProps/core.xml"] = xml_bytes(root)
            data = packed(parts)
            warnings.append("Only core document properties were cleared. This does not sanitize comments, revisions, custom properties or document contents.")
        elif slug == "word-page-size":
            sizes = {"A4": (210, 297), "Letter": (215.9, 279.4), "Legal": (215.9, 355.6)}
            size, orientation = options.get("size", "A4"), options.get("orientation", "portrait")
            if size not in sizes or orientation not in {"portrait", "landscape"}:
                raise ValueError("Choose a supported paper size and orientation.")
            width, height = sizes[size]
            if orientation == "landscape":
                width, height = height, width
            from docx.enum.section import WD_ORIENT
            for section in doc.sections:
                section.orientation = WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
                section.page_width, section.page_height = Mm(width), Mm(height)
        elif slug == "word-margins":
            margins = {key: Mm(bound(options, key, 25, 0, 100)) for key in ("top", "bottom", "left", "right")}
            for section in doc.sections:
                if margins["left"] + margins["right"] >= section.page_width - Mm(10) or margins["top"] + margins["bottom"] >= section.page_height - Mm(10):
                    raise ValueError("Margins must leave at least 10 mm of usable page width and height.")
                for key, value in margins.items():
                    setattr(section, f"{key}_margin", value)
        elif slug == "word-font":
            family = options.get("font", "Arial").strip()
            if not family or len(family) > 100 or re.search(r"[\x00-\x1f]", family):
                raise ValueError("Enter a valid font family (up to 100 characters).")
            size = bound(options, "fontSize", 12, 1, 200)
            from docx.text.run import Run
            for paragraph in all_paragraphs(doc):
                for element in paragraph._p.xpath(".//w:r"):
                    if any(a.tag == qn("w:txbxContent") for a in element.iterancestors()):
                        continue
                    run = Run(element, paragraph)
                    run.font.name, run.font.size = family, Pt(size)
                    fonts = run._element.get_or_add_rPr().rFonts
                    fonts.set(qn("w:eastAsia"), family)
                    fonts.set(qn("w:cs"), family)
        elif slug == "word-line-spacing":
            spacing = bound(options, "spacing", 1.5, 0.5, 3)
            for paragraph in body_paragraphs(doc):
                paragraph.paragraph_format.line_spacing = spacing
        elif slug == "word-header-footer":
            for section in doc.sections:
                for key, region in (("header", section.header), ("footer", section.footer)):
                    region.is_linked_to_previous = False
                    for node in list(region._element):
                        region._element.remove(node)
                    region.add_paragraph(options.get(key, ""))
            warnings.append("Default headers and footers were replaced, including any existing logos or fields. First/even-page variants were not changed.")
        elif slug == "remove-word-empty-paragraphs":
            count = 0
            for p in list(doc.element.body):
                if p.tag != qn("w:p") or "".join(p.xpath(".//w:t/text()")).strip():
                    continue
                # Do not remove fields, bookmarks, revision content, drawings or breaks.
                allowed = {qn("w:p"), qn("w:pPr"), qn("w:r"), qn("w:rPr"), qn("w:t")}
                protected = any(n.tag not in allowed and not (n.getparent() is not None and n.getparent().tag in {qn("w:pPr"), qn("w:rPr")})
                                for n in p.iter())
                if not protected and not p.xpath(".//w:sectPr"):
                    doc.element.body.remove(p)
                    count += 1
            message = f"Removed {count} empty body paragraphs. Check the resulting spacing."
        if data is None:
            data = document_bytes(doc)
    if extension == ".docx":
        # Reopen actual output, not merely trust a successful save.
        Document(io.BytesIO(data))
    files.append(output(slug + extension, data))
    return {"files": files, "message": message, "warnings": warnings, "text": text,
            "originalBytes": original, "outputBytes": len(data), "pageCount": 0}